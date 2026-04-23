import os
import re
import sys
import subprocess
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches
except ImportError as exc:
    raise SystemExit(
        "python-docx is required. Install with: pip install python-docx"
    ) from exc


def natural_sort_key(path_obj: Path):
    parts = re.split(r"(\d+)", path_obj.stem)
    key = []
    for part in parts:
        if part.isdigit():
            key.append(int(part))
        else:
            key.append(part.lower())
    return key


def extract_rmse_block(stdout_text: str) -> str:
    start_tag = "==================== RMSE summary ===================="
    end_tag = "======================================================"

    start_idx = stdout_text.find(start_tag)
    if start_idx < 0:
        return "RMSE summary not found in stdout."

    end_idx = stdout_text.find(end_tag, start_idx + len(start_tag))
    if end_idx < 0:
        return stdout_text[start_idx:].strip()

    end_idx = end_idx + len(end_tag)
    return stdout_text[start_idx:end_idx].strip()


def main():
    script_dir = Path(__file__).resolve().parent
    # data_dir = script_dir / "task_simulation_9000"
    # data_dir = script_dir / "task_simulation_50000"
    data_dir = script_dir / "task_simulation_90000"


    if not data_dir.exists():
        raise SystemExit(f"Data folder not found: {data_dir}")

    mat_files = sorted(data_dir.glob("*.mat"), key=natural_sort_key)
    if not mat_files:
        raise SystemExit(f"No .mat files found in: {data_dir}")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_root = script_dir / "batch_doc_output" / f"run_{timestamp}"
    images_dir = output_root / "images"
    output_root.mkdir(parents=True, exist_ok=True)
    images_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("ESKF Batch Simulation Report", level=1)
    doc.add_paragraph(f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"Data folder: {data_dir}")

    failures = []

    for idx, mat_file in enumerate(mat_files, start=1):
        case_name = mat_file.stem
        print(f"[{idx}/{len(mat_files)}] Running: {mat_file.name}")

        env = os.environ.copy()
        env["SIM_DATA_FILE"] = str(mat_file)
        env["SAVE_FIG_DIR"] = str(images_dir)
        env["FIG_PREFIX"] = case_name
        env["NO_SHOW_FIG"] = "1"
        env["MPLBACKEND"] = "Agg"

        cmd = [sys.executable, "run_INS_simulated.py"]
        result = subprocess.run(
            cmd,
            cwd=script_dir,
            env=env,
            text=True,
            capture_output=True,
        )

        doc.add_heading(case_name, level=2)
        doc.add_paragraph(f"Input data: {mat_file}")

        if result.returncode != 0:
            failures.append((mat_file.name, result.stderr.strip()))
            doc.add_paragraph("Run status: FAILED")
            doc.add_paragraph(result.stderr.strip() or "No stderr captured.")
            continue

        doc.add_paragraph("Run status: OK")
        doc.add_paragraph(extract_rmse_block(result.stdout))

        expected_images = [
            images_dir / f"{case_name}_fig1_trajectory.png",
            images_dir / f"{case_name}_fig2_states.png",
            images_dir / f"{case_name}_fig3_rmse.png",
            images_dir / f"{case_name}_fig4_rdiag.png",
        ]

        for image_path in expected_images:
            if image_path.exists():
                doc.add_picture(str(image_path), width=Inches(6.3))
            else:
                doc.add_paragraph(f"Missing image: {image_path.name}")

    if failures:
        doc.add_heading("Failures", level=2)
        for filename, err in failures:
            doc.add_paragraph(f"{filename}: {err[:1000]}")

    doc_path = output_root / "ESKF_batch_report.docx"
    doc.save(doc_path)

    print("\nBatch finished.")
    print(f"Report: {doc_path}")
    print(f"Images: {images_dir}")
    print(f"Failures: {len(failures)}")


if __name__ == "__main__":
    main()
